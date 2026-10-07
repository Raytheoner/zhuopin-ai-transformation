from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]
    style.font.name='Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.18
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(14)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 1 Char','Heading 2 Char','Heading 3 Char']:
    style=doc.styles[name]
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if 'theme' in key.lower(): del fonts.attrib[key]
    for scope in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+scope),'Microsoft YaHei')
    if style.font.size:
        size=style.element.get_or_add_rPr().find(qn('w:szCs'))
        if size is None: size=OxmlElement('w:szCs');style.element.get_or_add_rPr().append(size)
        size.set(qn('w:val'),str(int(style.font.size.pt*2)))
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
md=[]
header=s.header.paragraphs[0]
header.text='卓品智能  智能门户UI设计审核稿  2026年10月7日'
header.style='Caption'
header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=s.footer.paragraphs[0]
footer.text='设计审核用  全部业务数据和操作为模拟  '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(text):
    if text: md.append(text)
    return doc.add_paragraph(text)
def h(text):
    md.append('## '+text)
    heading=doc.add_heading(text,1)
    for run in heading.runs:
        run.font.name='Microsoft YaHei';run.font.size=Pt(19);run.bold=False
        run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    return heading
def image(file,width=7.1):
    md.append('!['+file+'](screens/'+file+')')
    doc.add_picture(str(ROOT/'screens'/file),width=Inches(width))
def table(headers,rows,widths):
    md.append('\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(str(v).replace('|','\\|') for v in row)+' |' for row in rows]))
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,(title,width) in enumerate(zip(headers,widths)):
        t.columns[i].width=Inches(width);t.rows[0].cells[i].text=title
    rpt=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rpt)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=1
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF8');cell._tc.get_or_add_tcPr().append(sh)
            for par in cell.paragraphs:
                par.paragraph_format.space_before=Pt(4);par.paragraph_format.space_after=Pt(4)
                for r in par.runs:r.font.size=Pt(10);r.bold=ri==0
    p('')
    return t
def new(text):h(text).paragraph_format.page_break_before=True

doc.add_paragraph('智能门户高保真UI方案审核稿','Title')
md.append('# 智能门户高保真UI方案审核稿')
p('供 Shao Peishen 审核整套业务体验。已选择方向1 明亮企业办公，修改意见无。本稿覆盖25类桌面页面、4类移动关键页面、统一UI规范与分模块应用建议。')
p('本轮完成设计与模拟原型，不改变生产API、认证、业务规则或部署。整套体验审核后另拟正式实施计划，合入和生产发布逐项授权。')
image('P02.jpg')
p('P02 门户首页  1440×1024  六部门加应用入口  全部数字与业务操作均为模拟。')
p('审阅入口 http://127.0.0.1:4173/  完整25页效果图见同目录效果图总览.html。预览依赖本机服务存续，原型源码包可复现。')

new('质量完整样板')
p('QD-B采用库内现有报告的六段结构和13模块。演示评分81不映射正式等级，委员会或PMO保留最终决定权。Q2采用D1到D7，既有分级、红线和质量工程师签发责任不因界面调整改变。')
image('P14.jpg')
p('P14 QD-B预审报告  先看建议，再看规则与证据，补齐缺失资料后进入人工复核。')
p('上传页继承XLSX及20 MB上限，模板为EQQR8082 A2.1。原型只校验文件名称、类型和大小，不读取正文或上传。C05与档三未签认的效力缺口保留。')

new('采购统一适配')
p('SC8覆盖成品、物料、案例和判例批改。四色同时提供文字说明，支持中文筛选、排序、分页、依据展开。物料使用独立样例数量，避免混用成品统计。')
image('P19.jpg')
p('P19 SC8物料覆盖  来源及更新时间就近展示，缺资料不等同无风险。')
p('SC2周报保留期次、来源和人工备注。现网SC2打开超时、SC8结果表未核验的缺口继续登记。缓存刷新与业务重算分开，本稿不新增齐套规则、R2阈值或真实六键取数。')

new('财务参考与调整边界')
p('财务需求回信尚未正式入档。FI2输入、结果和FI3付款清单作为信息结构参考，均显示需求待确认，未冻结默认查询、差异阈值、结案流程或付款权限。')
image('P23.jpg')
p('P23 FI2对账输入  本原型展示三源参考结构，数据导入及查询均为模拟。')
p('FI3逐项勾选及记录只产生演示反馈。收款方、账户尾号和金额为模拟，付款状态保持未执行。财务回件归档后再确认字段与适配顺序。')

new('移动端关键流程')
p('移动端重点覆盖部门与应用入口、待办、结果查看和人工复核。底部导航显示当前项，侧栏关闭后不进入键盘焦点，支持关闭按钮和Escape返回触发控件。')
t=doc.add_table(rows=1,cols=2);t.autofit=False
for c,f,title in zip(t.rows[0].cells,['M01.jpg','M02.jpg'],['M01 部门与应用入口','M02 我的待办']):
    md.append('!['+title+'](screens/'+f+')')
    c.width=Inches(3.55);c.paragraphs[0].add_run(title).bold=True;c.add_paragraph().add_run().add_picture(str(ROOT/'screens'/f),width=Inches(2.8))
p('390×844首屏样例。长页面正常滚动，宽表格限于表格容器横向查看，页面本身没有横向溢出。企微实际WebView与角色映射留待项目实施验收。')

new('移动结果与人工复核')
p('结果按建议、依据、缺失资料分层。人工决定必选，意见至少6字的表单校验为设计建议。反馈明确仅模拟记录，没有正式签认或生产写入。')
t=doc.add_table(rows=1,cols=2);t.autofit=False
for c,f,title in zip(t.rows[0].cells,['M03.jpg','M04-feedback.jpg'],['M03 结果查看','M04 模拟复核反馈']):
    md.append('!['+title+'](screens/'+f+')')
    c.width=Inches(3.55);c.paragraphs[0].add_run(title).bold=True;c.add_paragraph().add_run().add_picture(str(ROOT/'screens'/f),width=Inches(2.8))
p('复核记录仅在本次预览的页面会话内存中保留，刷新即清空。生产实施须绑定正式身份、权限与审计接口，L2不代签。')

new('统一UI规范')
table(['变量','规范'],[
['颜色','品牌蓝#1677ff；主按钮#126bd9；主文字#18253b；辅助文字#63738b；背景#f5f8fc；白色业务卡片。'],
['字体','Microsoft YaHei等本地字体。业务标题28px、正文14px、元信息12px；移动正文14px、元信息11px。'],
['布局','桌面1440×1024，侧栏214px，顶栏66px；移动390×844，内容边距17px，底部四项导航。'],
['组件','导航、检索、筛选、表格、上传、报告、证据、人工复核、来源时间与反馈采用共同表达。'],
['资产','Heroicons标准图标及独立生成插画。蓝色品牌图形为概念占位，应用前替换已核的官方资产。'],
['键盘','原生Tab、Enter与空格，跳到主要内容、可见焦点及减少运动；移动底部当前项提供aria-current。']], [1,6.1])
p('颜色不独立承载状态。服务在线、AI建议或绿色状态都不能等同业务通过。所有来源和更新时间就近可见，mock、试点、规划与需求待确认保持明确。')
table(['状态','表达及动作'],[['加载','保留上下文，正在加载演示快照。'],['空数据','暂无符合条件记录，可调整筛选。'],['失败','演示加载失败，可模拟重试。'],['陈旧与缺资料','显示旧时间或缺口，先补资料再人工确认。'],['权限不足','提示核验角色，不默认扩大权限。'],['规划与试点','入口显式标阶段，规划没有真实执行按钮。']], [1.2,5.9])

new('公共门户与部门页面目录')
mapping=json.loads((ROOT/'需求页面映射.json').read_text(encoding='utf-8'))
table(['ID','页面','确认程度与主要边界'],[(v['id'],v['title'],v['status']+'；'+v['interaction']) for v in mapping[:12]],[.5,1.75,4.85])
p('确认程度区分业务边界与新的设计表达。统一待办、登录展示及系统状态聚合属于设计目标，当前尚无全部整合接口。销售、运营、研发采用明确的规划模板。')

new('质量采购与财务页面目录')
table(['ID','页面','确认程度与主要边界'],[(v['id'],v['title'],v['status']+'；'+v['interaction']) for v in mapping[12:]],[.5,1.75,4.85])
p('完整路由及来源映射见需求页面映射.json和页面索引.csv。规则效力仍以场景正本及后续签认记录为准，本稿不将早期PRD或界面样例追认为当前规则。')

new('验证结论与应用顺序')
p('原型完成25桌面与4移动页面留图，核验导航、中文检索、筛选排序分页、证据展开、缺资料校验、模拟上传进度、人工意见及反馈、清单校验、移动导航与键盘焦点。320及1024窄屏检查无整页横向溢出。')
p('四项Node定向逻辑测试及Vite构建通过。一次Luna只读复核提出的两项移动导航问题已修复并回验。没有运行项目矩阵或业务fixture，也不证明新引擎、机器人监听、队列迁移或生产全链通过。')
p('应用顺序建议为公共框架、QD-B、Q2、采购，财务回件确认后适配。继续保留一个门户入口和独立业务服务，不为统一界面而合并后台。每模块另立intent、design与tasks，定向验证及review后申请发布。')
review_heading=doc.add_heading('整套体验审核记录',2)
for run in review_heading.runs:
    run.font.name='Microsoft YaHei';run.font.size=Pt(14);run.bold=False
    run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
md.append('### 整套体验审核记录')
for text in [
'1 选择题  A通过进入正式实施计划  B修改后再审  C暂缓。答复不构成部署批准。',
'2 判断题  六部门加应用入口的首页结构符合工作习惯。是 / 否。',
'3 判断题  建议、依据、缺失资料和人工决定分层清楚。是 / 否。',
'4 选择题  A按上述顺序实施  B调整顺序并说明。',
'5 请按页面ID填写需要修改的字段、信息或操作。',
'6 请提供已核官方品牌图形的仓内路径，或在实施前另行确认资产。']:
    p(text)
p('视觉方向1已经选定，无需重复回答。本稿的整套业务体验尚待审核，正式实施及生产发布保留后续逐项授权。')
doc.save(ROOT/'智能门户UI设计审核稿.docx')
(ROOT/'智能门户UI设计审核稿.md').write_text('\n\n'.join(md)+'\n',encoding='utf-8')
print('DOCX draft created',ROOT/'智能门户UI设计审核稿.docx')
