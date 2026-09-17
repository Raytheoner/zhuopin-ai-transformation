"""把唐燕萍 2026-08-22 交付的节假日日历 xlsx 落成 FI3 可读的 CSV 数据件。

用法（xlsx 原件在 `7-外部文档/财务部/`，该目录不入库、只在主 checkout 存证）::

    python scripts/build_holiday_calendar.py "<xlsx 路径>"

输出 ``data/holidays/holiday_calendar.csv``，五列原样搬运（日期／星期／节假日名称／日期类型／
是否工作日），日期规整为 ``YYYY-MM-DD``。🔴 **不做任何推导、不补行、不删行**——本脚本只是搬运，
「工作日」判定直查 ``是否工作日`` 列（就绪包 §二.2 第 2 条），推导逻辑一律不在此出现。

落盘后打印行数／类型分布／覆盖区间／源文件 md5，供与就绪包 §一 的实测值逐项对表。
"""
from __future__ import annotations

import csv
import hashlib
import sys
from collections import Counter
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "data" / "holidays" / "holiday_calendar.csv"
COLUMNS = ("日期", "星期", "节假日名称", "日期类型", "是否工作日")


def main(xlsx: str) -> int:
    src = Path(xlsx)
    wb = openpyxl.load_workbook(src, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = tuple(rows[0])
    if header != COLUMNS:
        raise SystemExit(f"xlsx 列头与预期不符：{header!r} ≠ {COLUMNS!r}")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNS)
        for d, weekday, name, kind, is_workday in rows[1:]:
            w.writerow([d.strftime("%Y-%m-%d"), weekday, name or "", kind, is_workday])
    body = rows[1:]
    print(f"源文件 md5={hashlib.md5(src.read_bytes()).hexdigest()}")
    print(f"落盘 {OUT}：{len(body)} 行，覆盖 {body[0][0]:%Y-%m-%d} ～ {body[-1][0]:%Y-%m-%d}")
    print(f"日期类型分布：{dict(Counter(r[3] for r in body))}")
    print(f"是否工作日分布：{dict(Counter(r[4] for r in body))}")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    sys.exit(main(sys.argv[1]))
