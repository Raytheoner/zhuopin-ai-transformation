"""候选实际模块的合成业务/audit验收，不接真实模型。"""
import hashlib,io,json,re,time
from pathlib import Path
import openpyxl
from pptx import Presentation
from zhuopin_platform.audit import AuditLogger
from q2_8d_verdict.document_evidence import DocumentScope
from q2_8d_verdict.document_upload import DocumentUploadService,JobRef
from q2_8d_verdict.webapp import create_app
from test_document_upload import FakeModel,wait_job

class ContentModel(FakeModel):
    def request(self,request,operation_dir):
        response=super().request(request,operation_dir)
        if 'SYNTHETIC_LOW' in request.prompt:
            response.payload['items'][0].update(score=0,reason='合成原文明确SYNTHETIC_LOW，完整阅读后本项未满足',remediation='补齐本项缺失要求')
        return response

def run_case(root,low):
    root.mkdir()
    prs=Presentation()
    for i in range(1,8):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_textbox(0,0,4000000,600000).text=f'D{i}: 合成输入 '+('SYNTHETIC_LOW' if low else 'SYNTHETIC_HIGH')
    rawio=io.BytesIO();prs.save(rawio);raw=rawio.getvalue()
    scope=DocumentScope('SYNTHETIC-A','T1','B',hashlib.sha256(raw).hexdigest())
    audit=AuditLogger.jsonl(root/'audit.jsonl');model=ContentModel();service=DocumentUploadService(root/'store',audit,model,scope)
    app=create_app(audit_path=root/'web.jsonl',real_upload_enabled=True,real_upload_service=service,real_upload_actor=lambda:'local:test',real_upload_access=lambda a,b:True)
    c=app.test_client();page=c.get('/quality/q2/review')
    token=re.search('name="csrf_token" value="([^"]+)"',page.text).group(1)
    r=c.post('/quality/q2/review',data={'csrf_token':token,'file':(io.BytesIO(raw),'synthetic.pptx')},headers={'Origin':'http://localhost'})
    assert r.status_code==303
    object_id=r.location.rsplit('/',1)[1];saved=wait_job(service,JobRef(object_id,'receiving'))
    assert saved.state=='complete',saved.error
    report=c.get(r.location);assert report.status_code==200 and object_id in report.text and saved.receipt['event_sha256'] in report.text
    download=c.get(r.location+'/excel');assert download.status_code==200 and 'attachment' in download.headers['Content-Disposition']
    workbook=openpyxl.load_workbook(io.BytesIO(download.data),data_only=True)
    total=workbook['评审汇总']['B13'].value
    assert total==saved.payload['advice']['total']
    assert len(audit.query_by(scenario='Q2',action='real_document_ai_advice'))==1
    chain=audit.verify_chain();assert chain.ok and chain.total==1 and not chain.error
    retry=c.post('/quality/q2/review',data={'csrf_token':token,'file':(io.BytesIO(raw),'synthetic.pptx')},headers={'Origin':'http://localhost'})
    assert retry.location==r.location and model.calls==1
    return total,object_id

def test_two_distinct_inputs_produce_traceable_different_results(tmp_path):
    high,first=run_case(tmp_path/'high',False);low,second=run_case(tmp_path/'low',True)
    assert high>low and first!=second
