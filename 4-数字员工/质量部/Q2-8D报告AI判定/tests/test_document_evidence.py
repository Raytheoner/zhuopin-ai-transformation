import hashlib,zipfile
from pathlib import Path
import pytest
from pptx import Presentation
from q2_8d_verdict.document_evidence import DocumentScope, EvidenceError, read_pptx_evidence

def fixture(tmp_path, texts=('D3: 第一页临时措施','D3: 第二页结果与负责人')):
    p=tmp_path/'input.pptx'; prs=Presentation()
    for text in texts:
        s=prs.slides.add_slide(prs.slide_layouts[6]); s.shapes.add_textbox(0,0,5000000,500000).text=text
    prs.save(p)
    return p,DocumentScope('SYNTHETIC-A','T1','B',hashlib.sha256(p.read_bytes()).hexdigest())

def test_continued_chapter_and_locations(tmp_path):
    p,scope=fixture(tmp_path); b=read_pptx_evidence(p,scope)
    assert '第一页临时措施' in b.sections['D3'] and '第二页结果与负责人' in b.sections['D3']
    assert {r.page for r in b.records if r.text}=={1,2}
    assert b.complete and b.paragraph_count==2

def test_table_multiline_and_empty_cells(tmp_path):
    prs=Presentation(); s=prs.slides.add_slide(prs.slide_layouts[6])
    t=s.shapes.add_table(2,2,0,0,5000000,500000).table
    t.cell(0,0).text='D3:措施'; t.cell(1,0).text='行一\n行二'; t.cell(1,1).text='=不能执行'
    p=tmp_path/'table.pptx'; prs.save(p)
    b=read_pptx_evidence(p,DocumentScope('SYNTHETIC-A','T1','B',hashlib.sha256(p.read_bytes()).hexdigest()))
    assert any(r.text=='行一\n行二' for r in b.records)
    assert any(r.kind=='empty_table_cell' for r in b.records)
    assert any(r.text=='=不能执行' for r in b.records)
    assert b.paragraph_count==4

def test_source_hash_and_asil_are_enforced(tmp_path):
    p,s=fixture(tmp_path)
    for scope in (DocumentScope(s.oem_label,s.source_version,'C',s.authorized_sha256),DocumentScope(s.oem_label,s.source_version,'B','0'*64)):
        with pytest.raises(EvidenceError):read_pptx_evidence(p,scope)

def test_bad_zip_not_empty_document(tmp_path):
    p=tmp_path/'bad.pptx';p.write_bytes(b'bad')
    with pytest.raises(EvidenceError):read_pptx_evidence(p,DocumentScope('A','T','B',hashlib.sha256(b'bad').hexdigest()))

def test_unsafe_archive_name_is_rejected(tmp_path):
    p=tmp_path/'bad.pptx'
    with zipfile.ZipFile(p,'w') as z:z.writestr('../outside',b'no')
    with pytest.raises(EvidenceError):read_pptx_evidence(p,DocumentScope('A','T','B',hashlib.sha256(p.read_bytes()).hexdigest()))

def test_asset_decode_failure_marks_incomplete(tmp_path):
    p,scope=fixture(tmp_path)
    with zipfile.ZipFile(p,'a') as z:z.writestr('ppt/media/broken.png',b'bad image')
    scope=DocumentScope('A','T','B',hashlib.sha256(p.read_bytes()).hexdigest())
    b=read_pptx_evidence(p,scope)
    assert not b.complete and b.warnings and b.assets[0]['read_state']=='READ_ERROR'
