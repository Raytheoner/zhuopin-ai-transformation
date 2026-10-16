from dataclasses import replace
import hashlib,zipfile
import openpyxl
from q2_8d_verdict.document_evidence import EvidenceRecord
from q2_8d_verdict.ai_advice import validate_advice
from q2_8d_verdict.excel_report import build_exports
from test_ai_advice import advice_fixture

def test_five_sheets_values_formulas_and_no_source_execution(tmp_path,ruleset):
    b,v,p=advice_fixture(ruleset)
    long='=原文不是公式\n'+'12345678901234567890中文'*2000
    b=replace(b,records=b.records+(EvidenceRecord('long',2,'2','text',long),))
    r=validate_advice(p,b,v,ruleset);out=build_exports(r,b,tmp_path/'export')
    values=openpyxl.load_workbook(out.xlsx,data_only=True);formulas=openpyxl.load_workbook(out.xlsx,data_only=False)
    assert values.sheetnames==['评审明细','评审汇总','原文依据','图件依据','附件数据']
    assert values['评审汇总']['B13'].value==r.total and values['评审汇总']['B14'].value==r.numeric_grade
    assert formulas['评审汇总']['B13'].data_type=='f'
    parts=[row[4].value for row in values['原文依据'] if row[0].value=='long']
    assert ''.join(parts)==long
    assert all(row[4].data_type!='f' for row in formulas['原文依据'] if row[0].value=='long')
    assert max(d.height or 15 for s in values for d in s.row_dimensions.values())<=409.5
    with zipfile.ZipFile(out.bundle) as z:assert 'manifest.json' in z.namelist()

def test_assets_cached_cells_and_zip_hashes(tmp_path,ruleset):
    from PIL import Image
    b,v,p=advice_fixture(ruleset);image=tmp_path/'image.png';Image.new('RGB',(200,100),'white').save(image)
    a={'ref':'asset-1','kind':'media','part':'ppt/media/x.png','path':str(image),'preview':str(image),'sha256':hashlib.sha256(image.read_bytes()).hexdigest(),'pages':[1],'read_state':'READ'}
    data=tmp_path/'data.bin';data.write_bytes(b'synthetic')
    table={'ref':'asset-2','kind':'workbook','part':'ppt/embeddings/a.bin','path':str(data),'sha256':hashlib.sha256(data.read_bytes()).hexdigest(),'pages':[1],'read_state':'READ','cells':[{'sheet':'S','row':1,'column':1,'value':'12345678901234567890','value_type':'text','cache_state':'READ'},{'sheet':'S','row':2,'column':1,'value':'=literal','value_type':'text','cache_state':'READ'}]}
    b=replace(b,assets=(a,table));r=validate_advice(p,b,v,ruleset);out=build_exports(r,b,tmp_path/'export')
    w=openpyxl.load_workbook(out.xlsx,data_only=False)
    assert len(w['图件依据']._images)==1
    assert w['附件数据']['G5'].value=='12345678901234567890' and w['附件数据']['G6'].value=='=literal' and w['附件数据']['G6'].data_type!='f'
    with zipfile.ZipFile(out.bundle) as z:
        assert hashlib.sha256(z.read('assets/'+image.name)).hexdigest()==a['sha256']
