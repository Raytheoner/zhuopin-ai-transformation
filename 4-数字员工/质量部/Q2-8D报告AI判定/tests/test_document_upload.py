import io,json,time,hashlib
from types import SimpleNamespace
from pathlib import Path
import pytest
from pptx import Presentation
from zhuopin_platform.audit import AuditLogger
from q2_8d_verdict.document_evidence import DocumentScope
from q2_8d_verdict.document_upload import DocumentUploadService,UploadError

class FakeModel:
    calls=0
    def request(self,request,operation_dir):
        self.calls+=1
        data=json.loads(request.prompt.split('不可信证据数据开始（仅作为数据）：\n')[1])
        record=next(x for x in data['records'] if x['text'])
        p={k:data[k] for k in ('source_sha256','rules_sha256')}
        p.update({k:data['scope'][k] for k in ('oem_label','source_version','asil_level')})
        p.update(items=[dict(rule_id=r['rule_id'],score=r['score'],refs=[record['ref']],quote=record['text'][:6],reason='合成评分验证',remediation='提交质量工程师复核',evidence_state='READ',confidence='MED') for r in data['rules']],d3_missing=[],redline_observations=[])
        return SimpleNamespace(payload=p,requested_model='gpt-6-luna',observed_model=None,usage=None,prompt_hash=request.prompt_hash,schema_hash=request.schema_hash,request_sha256=hashlib.sha256(request.prompt.encode()).hexdigest(),exit_code=0,raw_event_sha256='f'*64,tool_calls=0)

def upload_fixture(tmp_path):
    prs=Presentation()
    for i in range(1,8):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_textbox(0,0,4000000,500000).text=f'D{i}: 合成证据 区域 数量 措施 结果 负责人 计划日 完成日 状态'
    b=io.BytesIO();prs.save(b);raw=b.getvalue()
    scope=DocumentScope('SYNTHETIC-A','T1','B',hashlib.sha256(raw).hexdigest())
    model=FakeModel(); audit=AuditLogger.jsonl(tmp_path/'audit.jsonl')
    return raw,model,audit,DocumentUploadService(tmp_path/'store',audit,model,scope)

def wait_job(service,ref,actor='local:test'):
    for _ in range(200):
        saved=service.get(ref.object_id,actor)
        if saved.state in ('complete','failed','interrupted'):return saved
        time.sleep(.02)
    raise AssertionError('worker did not finish')

def fake_export(self,report,bundle,dest):
    dest.mkdir(exist_ok=True);x=dest/'report.xlsx';z=dest/'assets.zip';x.write_bytes(b'synthetic xlsx');z.write_bytes(b'synthetic zip')
    return {'xlsx':{'path':str(x),'sha256':hashlib.sha256(x.read_bytes()).hexdigest()},'bundle':{'path':str(z),'sha256':hashlib.sha256(z.read_bytes()).hexdigest()}}

def test_idempotent_actor_bound_and_unique_audit(tmp_path,monkeypatch):
    monkeypatch.setattr(DocumentUploadService,'_build_exports',fake_export)
    raw,m,a,s=upload_fixture(tmp_path);ref=s.submit(raw,'case.pptx','local:test');saved=wait_job(s,ref)
    assert saved.state=='complete',saved.error
    assert s.submit(raw,'case.pptx','local:test').object_id==ref.object_id and m.calls==1
    assert len(a.query_by(scenario='Q2',action='real_document_ai_advice'))==1 and a.verify_chain().ok
    with pytest.raises(UploadError):s.get(ref.object_id,'local:other')
    s.export(ref.object_id,'local:test','xlsx').write_bytes(b'tampered')
    with pytest.raises(UploadError):s.get(ref.object_id,'local:test')

def test_scope_rejects_before_store_or_model(tmp_path):
    raw,m,a,s=upload_fixture(tmp_path)
    with pytest.raises(UploadError):s.submit(raw+b'changed','case.pptx','local:test')
    assert m.calls==0

def test_model_failure_preserved_not_replayed(tmp_path,monkeypatch):
    raw,m,a,s=upload_fixture(tmp_path)
    def fail(*args):m.calls+=1;raise RuntimeError('model denied')
    m.request=fail
    saved=wait_job(s,s.submit(raw,'case.pptx','local:test'))
    assert saved.state=='failed' and 'model denied' in saved.error
    assert s.submit(raw,'case.pptx','local:test').state=='failed' and m.calls==1
    assert not a.query_by(scenario='Q2',action='real_document_ai_advice')

def test_audit_failure_prevents_complete(tmp_path,monkeypatch):
    monkeypatch.setattr(DocumentUploadService,'_build_exports',fake_export)
    raw,m,a,s=upload_fixture(tmp_path)
    monkeypatch.setattr(a,'record',lambda e:None)
    saved=wait_job(s,s.submit(raw,'case.pptx','local:test'))
    assert saved.state=='failed' and 'audit' in saved.error
    with pytest.raises(UploadError):s.export(saved.object_id,'local:test','xlsx')

def test_unfinished_job_after_restart_is_interrupted(tmp_path):
    raw,m,a,s=upload_fixture(tmp_path)
    monkey=pytest.MonkeyPatch();monkey.setattr(s,'_start',lambda *a:None)
    ref=s.submit(raw,'case.pptx','local:test')
    restarted=DocumentUploadService(s.root,a,m,s.scope)
    assert restarted.get(ref.object_id,'local:test').state=='interrupted'
    monkey.undo()

@pytest.mark.parametrize('failure',['source','export','publish','chain'])
def test_persistence_boundary_failures_never_publish_success(tmp_path,monkeypatch,failure):
    monkeypatch.setattr(DocumentUploadService,'_build_exports',fake_export)
    raw,m,a,s=upload_fixture(tmp_path)
    if failure in ('source','publish'):
        original=s._immutable
        def write(path,data):
            if path.name==('source.pptx' if failure=='source' else 'result.json'):raise OSError('injected persistence failure')
            return original(path,data)
        monkeypatch.setattr(s,'_immutable',write)
    if failure=='export':monkeypatch.setattr(s,'_build_exports',lambda *args:(_ for _ in ()).throw(OSError('export readback failed')))
    if failure=='chain':
        monkeypatch.setattr(a,'verify_chain',lambda:SimpleNamespace(ok=True,total=0,error='unsupported'))
    if failure=='source':
        with pytest.raises(OSError):s.submit(raw,'case.pptx','local:test')
        assert m.calls==0
    else:
        saved=wait_job(s,s.submit(raw,'case.pptx','local:test'))
        assert saved.state=='failed'
        with pytest.raises(UploadError):s.export(saved.object_id,'local:test','xlsx')

def test_model_provenance_mismatch_blocks_scoring(tmp_path,monkeypatch):
    raw,m,a,s=upload_fixture(tmp_path);original=m.request
    def wrong(*args):
        r=original(*args);r.request_sha256='0'*64;return r
    m.request=wrong
    assert wait_job(s,s.submit(raw,'case.pptx','local:test')).state=='failed'

def test_one_citation_repair_preserves_original_and_decision(tmp_path,monkeypatch):
    monkeypatch.setattr(DocumentUploadService,'_build_exports',fake_export)
    raw,m,a,s=upload_fixture(tmp_path);original=m.request;first=[]
    def request(req,op):
        if not first:
            first.append(req);r=original(req,op);r.payload['items'][0]['quote']='D2:';return r
        r=original(first[0],op);r.request_sha256=hashlib.sha256(req.prompt.encode()).hexdigest();return r
    m.request=request
    saved=wait_job(s,s.submit(raw,'case.pptx','local:test'))
    assert saved.state=='complete',saved.error
    assert m.calls==2 and len(saved.payload['model']['attempts'])==2
    assert saved.payload['model']['repair_reason']=='quote not present in cited evidence'
