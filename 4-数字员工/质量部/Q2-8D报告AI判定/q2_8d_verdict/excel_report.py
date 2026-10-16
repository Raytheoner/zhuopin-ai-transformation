"""可独立查阅的五表 Excel；完整原文/缓存值留存，来源文本不执行。"""
from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
import hashlib,json,math,zipfile
import openpyxl
import xlsxwriter
from . import config
from .ai_advice import canonical

SHEETS=('评审明细','评审汇总','原文依据','图件依据','附件数据')

@dataclass(frozen=True)
class ExportFiles:
    xlsx:Path
    bundle:Path
    def to_dict(self):return {k:{'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for k,p in (('xlsx',self.xlsx),('bundle',self.bundle))}

def _chunks(value,size=300):
    text=str(value)
    return [text[i:i+size] for i in range(0,len(text),size)] or ['']

def build_exports(report,bundle,target_dir:Path)->ExportFiles:
    if not bundle.complete or report.professional_signoff:raise ValueError('only complete unconfirmed advice may export')
    target_dir.mkdir(parents=True,exist_ok=False)
    xlsx=target_dir/'8D完整AI建议评审报告.xlsx';archive=target_dir/'8D原资产与评分.zip'
    w=xlsxwriter.Workbook(xlsx,{'strings_to_formulas':False,'strings_to_urls':False,'nan_inf_to_errors':False})
    w.set_properties({'title':'8D完整AI建议评审报告','subject':'未签认AI建议，质量工程师最终复核','author':'Q2业务试评'})
    fmt=w.add_format({'font_name':'微软雅黑','font_size':10,'text_wrap':True,'valign':'top','border':1,'border_color':'#E0E6ED'})
    number=w.add_format({'font_name':'微软雅黑','font_size':10,'num_format':'0.##','border':1,'border_color':'#E0E6ED','valign':'top'})
    title=w.add_format({'font_name':'微软雅黑','font_size':19,'bold':True,'font_color':'#17365D'})
    banner=w.add_format({'font_name':'微软雅黑','font_size':11,'bold':True,'font_color':'#9C5700','bg_color':'#FFF2CC','text_wrap':True,'valign':'top'})
    head=w.add_format({'font_name':'微软雅黑','font_size':10,'bold':True,'font_color':'white','bg_color':'#24476B','text_wrap':True,'valign':'vcenter'})
    link=w.add_format({'font_name':'微软雅黑','font_size':10,'font_color':'#0563C1','underline':1,'text_wrap':True,'valign':'top','border':1})
    sheets={name:w.add_worksheet(name) for name in SHEETS}
    def init(name,headers,widths):
        s=sheets[name];s.hide_gridlines(2);s.freeze_panes(4,2);s.set_zoom(85)
        s.merge_range(0,0,0,len(headers)-1,name+' · Q2单件8D建议评审',title);s.set_row(0,32)
        s.merge_range(1,0,1,len(headers)-1,'AI建议 · 未专业签认 · 数字等级不代表业务验收通过 · 最终由质量工程师复核/签发',banner);s.set_row(1,35)
        s.merge_range(2,0,2,len(headers)-1,f"{bundle.scope.oem_label} / {bundle.scope.source_version} / ASIL {bundle.scope.asil_level} | 场景：{report.scene}（{'、'.join(report.scene_flags) or '已标注'}） | 原件SHA256：{bundle.source_sha256}",fmt);s.set_row(2,32)
        for i,(h,width) in enumerate(zip(headers,widths)):s.write_string(3,i,h,head);s.set_column(i,i,width)
        s.set_row(3,32);s.set_landscape();s.set_paper(8);s.fit_to_pages(1,0);s.repeat_rows(0,3)
        return s
    def text(s,row,col,value,format=fmt):s.write_string(row,col,str(value),format)
    def row_height(s,row,parts,widths):
        lines=max((sum(max(1,math.ceil(sum(2 if ord(c)>255 else 1 for c in line)/width)) for line in str(t).split('\n')) for t,width in zip(parts,widths)),default=1)
        s.set_row(row,min(409.5,max(26,lines*15+8)))
    source=init('原文依据',['证据定位','页码','形状/单元格','分片顺序','完整原文（连续分片，不截断）','读取类型'],[34,8,24,12,100,20]);refrows={};row=4
    for record in bundle.records:
        chunks=_chunks(record.text);refrows[record.ref]=('原文依据',row)
        for i,part in enumerate(chunks,1):
            vals=[record.ref,record.page,record.shape,f'{i}/{len(chunks)}',part,record.kind]
            for col,value in enumerate(vals):text(source,row,col,value)
            row_height(source,row,[part],[100]);row+=1
    source.autofilter(3,0,max(3,row-1),5)
    figure=init('图件依据',['图件定位','页码','原包路径','读取状态','图件/原始hash','说明'],[34,10,38,18,70,32]);row=4
    manifest=[]
    for asset in bundle.assets:
        manifest.append({k:v for k,v in asset.items() if k not in ('cells','path','preview')})
        if asset['kind']!='media':continue
        refrows[asset['ref']]=('图件依据',row)
        for col,val in enumerate((asset['ref'],','.join(map(str,asset['pages'])),asset['part'],asset['read_state'],asset['sha256'],'原二进制在资产ZIP；本表为等比预览')):text(figure,row,col,val)
        figure.set_row(row,44);row+=1
        from PIL import Image
        with Image.open(asset['preview']) as im:width,height=im.size
        scale=min(620/width,260/height,1)
        figure.insert_image(row,2,asset['preview'],{'x_scale':scale,'y_scale':scale,'object_position':1})
        figure.set_row(row,205);row+=2
    cached=init('附件数据',['附件定位','原包路径','工作表','行','列','类型/状态','已保存值（不执行公式）','分片顺序'],[28,40,22,8,8,20,75,12]);row=4
    for asset in bundle.assets:
        if asset['kind']!='workbook':continue
        refrows[asset['ref']]=('附件数据',row)
        for cell in asset.get('cells',[]):
            ref=f"{asset['ref']}:{cell['sheet']}!R{cell['row']}C{cell['column']}";refrows[ref]=('附件数据',row)
            parts=_chunks(cell['value']) if cell['value_type']=='text' else [cell['value']]
            for i,part in enumerate(parts,1):
                vals=[ref,asset['part'],cell['sheet'],cell['row'],cell['column'],cell['value_type']+'/'+cell['cache_state'],part,f'{i}/{len(parts)}']
                for col,value in enumerate(vals):
                    if col==6 and isinstance(value,(int,float)) and not isinstance(value,bool):cached.write_number(row,col,value,number)
                    else:text(cached,row,col,value)
                row_height(cached,row,[str(part)],[75]);row+=1
    cached.autofilter(3,0,max(3,row-1),7)
    detail=init('评审明细',['规则ID','章节','评分维度','满分','AI建议分','原基线分','差异/状态','原文依据','得扣分理由','整改建议','原文定位','规则判据'],[13,8,23,8,10,10,22,55,55,55,35,45]);row=4
    for item in report.items:
        fields=[item['rule_id'],item['step'],item['dimension'],item['max_score'],item['score'],item['baseline_score'],item['baseline_status']+' → AI建议'+('（有差异）' if item['changed'] else ''),item['quote'],item['reason'],item['remediation'],'\n'.join(item['refs']),item['criteria']]
        chunks=[_chunks(x) for x in fields];count=max(map(len,chunks))
        for offset in range(count):
            vals=[p[offset] if offset<len(p) else '' for p in chunks]
            for col,value in enumerate(vals):
                if col in (3,4,5) and offset==0:detail.write_number(row,col,float(fields[col]),number)
                elif col in (0,1) and offset>0:text(detail,row,col,fields[col])
                elif col==10 and offset==0:
                    sheet,anchor=refrows[item['refs'][0]]
                    detail.write_url(row,col,f"internal:'{sheet}'!A{anchor+1}",link,value)
                else:text(detail,row,col,value)
            row_height(detail,row,vals,[13,8,23,8,10,10,22,55,55,55,35,45]);row+=1
    last=row
    detail.autofilter(3,0,last-1,11)
    summary=init('评审汇总',['指标','AI建议得分/状态','适用满分','说明'],[34,30,18,110])
    for index,step in enumerate(sorted(report.chapter_scores),4):
        text(summary,index,0,step)
        expression=f'SUMIF(\'评审明细\'!B5:B{last},A{index+1},\'评审明细\'!E5:E{last})'
        if step=='D3':expression+=f'-{report.d3_deduction}'
        summary.write_formula(index,1,'='+expression,number,report.chapter_scores[step]);summary.write_number(index,2,report.chapter_max[step],number)
        text(summary,index,3,'章分，AI建议，待专业复核')
    text(summary,11,0,'D3额外扣分（已计入章分）');summary.write_number(11,1,-report.d3_deduction,number);text(summary,11,3,'；'.join(report.d3_missing) or '未列缺失字段')
    text(summary,12,0,'AI建议总分');summary.write_formula(12,1,'=MAX(0,SUM(B5:B11))',number,report.total);summary.write_number(12,2,report.max_score,number)
    text(summary,12,3,'数字评分用于修改参考，不代表正式业务放行')
    text(summary,13,0,'数字建议等级')
    grade='"'+config.CRITERIA.value_of('M5_GRADE_BANDS')[-1][0]+'"'
    for name,floor in reversed(config.CRITERIA.value_of('M5_GRADE_BANDS')[:-1]):grade=f'IF(B13>={floor},"{name}",{grade})'
    summary.write_formula(13,1,'='+grade,fmt,report.numeric_grade);text(summary,13,3,'分级阈值引用原签认M5；红线与处置单独判断')
    for index,values in enumerate([('处置','转人工裁决','L2；质量工程师最终复核/签发'),('专业签认','未签认','本Excel不代签、不自动退回/归档'),('读取完整性',bundle.read_state,f'{bundle.page_count}页；{bundle.paragraph_count}非空段；{len(bundle.assets)}原资产'),('规则版本',report.rule_version,report.rules_sha256),('提示词版本/Hash',report.prompt_hash,'schema '+report.schema_hash)],15):
        text(summary,index,0,values[0]);text(summary,index,1,values[1]);text(summary,index,3,values[2]);summary.set_row(index,42)
    row=21
    for redline in report.redlines:
        text(summary,row,0,'红线'+str(redline['no']));text(summary,row,1,redline['status']);text(summary,row,3,redline['description']+'；'+canonical(redline['observation']));summary.set_row(row,70);row+=1
    w.close()
    # 独立读回实际文件；不是检查生成前的 JSON。
    values=openpyxl.load_workbook(xlsx,data_only=True);formulas=openpyxl.load_workbook(xlsx,data_only=False)
    try:
        if values.sheetnames!=list(SHEETS) or values['评审汇总']['B13'].value!=report.total or values['评审汇总']['B14'].value!=report.numeric_grade:raise ValueError('Excel score readback failed')
        rebuilt={}
        for cells in values['原文依据'].iter_rows(min_row=5,values_only=True):rebuilt.setdefault(cells[0],[]).append(cells[4] or '')
        if any(''.join(rebuilt.get(r.ref,[]))!=r.text for r in bundle.records):raise ValueError('Excel source text coverage failed')
        for sf,sv in zip(formulas,values):
            for cells in sf:
                for cell in cells:
                    if cell.data_type=='f' and (sv[cell.coordinate].value is None or sv[cell.coordinate].data_type=='e'):raise ValueError('Excel formula cache failed')
    finally:values.close();formulas.close()
    with zipfile.ZipFile(archive,'x',compression=zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx,xlsx.name);included=set()
        for a in bundle.assets:
            for key in ('path','preview'):
                if not a.get(key):continue
                path=Path(a[key]);name='assets/'+path.name
                if name not in included:z.write(path,name);included.add(name)
        z.writestr('manifest.json',canonical({'scope':report.scope,'source_sha256':report.source_sha256,'assets':manifest,'xlsx_sha256':hashlib.sha256(xlsx.read_bytes()).hexdigest(),'professional_signoff':False}))
    with zipfile.ZipFile(archive) as z:
        if z.testzip() is not None:raise ValueError('asset ZIP readback failed')
        for asset in bundle.assets:
            if hashlib.sha256(z.read('assets/'+Path(asset['path']).name)).hexdigest()!=asset['sha256']:raise ValueError('asset hash mismatch')
    return ExportFiles(xlsx,archive)
