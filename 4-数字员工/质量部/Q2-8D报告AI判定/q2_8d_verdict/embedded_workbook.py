"""只读嵌入工作簿的已保存值；不运行公式、宏、外部链接。"""
from __future__ import annotations
from dataclasses import dataclass
from io import BytesIO
import zipfile


class EmbeddedReadError(ValueError):
    pass


@dataclass(frozen=True)
class CachedCell:
    sheet: str
    row: int
    column: int
    value: object
    value_type: str
    cache_state: str = "READ"


def read_cached_cells(raw: bytes, name: str) -> tuple[CachedCell, ...]:
    try:
        if raw.startswith(b"PK"):
            import openpyxl
            with zipfile.ZipFile(BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
                    raise EmbeddedReadError("embedded workbook too large")
            values = openpyxl.load_workbook(BytesIO(raw), data_only=True, read_only=True, keep_links=False)
            formulas = openpyxl.load_workbook(BytesIO(raw), data_only=False, read_only=True, keep_links=False)
            out = []
            try:
                for sheet in formulas:
                    if sheet.max_row * sheet.max_column > 500000:
                        raise EmbeddedReadError("embedded cell limit")
                    cached = values[sheet.title]
                    for row in sheet:
                        for cell in row:
                            if not hasattr(cell, 'row'):
                                continue
                            value = cached.cell(cell.row, cell.column).value
                            if cell.data_type == "f" and value is None:
                                raise EmbeddedReadError(f"formula cache missing: {sheet.title}!{cell.coordinate}")
                            # inlineStr with an empty value is a saved empty string, not an absent cell.
                            if cell.value is None and cell.data_type not in ("s", "inlineStr"):
                                continue
                            if value is None:
                                value = ""
                            typ = "text" if isinstance(value, str) else "boolean" if isinstance(value, bool) else "number"
                            out.append(CachedCell(sheet.title, cell.row, cell.column, value, typ))
            finally:
                values.close(); formulas.close()
            return tuple(out)
        if raw.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            import xlrd
            book = xlrd.open_workbook(file_contents=raw, on_demand=True)
            out = []
            try:
                for sheet in book.sheets():
                    if sheet.nrows * sheet.ncols > 500000:
                        raise EmbeddedReadError("embedded cell limit")
                    for row in range(sheet.nrows):
                        for col in range(sheet.ncols):
                            cell = sheet.cell(row, col)
                            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                                continue
                            if cell.ctype == xlrd.XL_CELL_ERROR:
                                raise EmbeddedReadError(f"cached cell error: {sheet.name}!R{row+1}C{col+1}")
                            typ = "text" if cell.ctype == xlrd.XL_CELL_TEXT else "boolean" if cell.ctype == xlrd.XL_CELL_BOOLEAN else "number"
                            out.append(CachedCell(sheet.name, row+1, col+1, cell.value, typ))
            finally:
                book.release_resources()
            return tuple(out)
        raise EmbeddedReadError(f"unsupported embedded format: {name}")
    except EmbeddedReadError:
        raise
    except Exception as exc:
        raise EmbeddedReadError(f"embedded read failed ({type(exc).__name__}): {name}") from exc
