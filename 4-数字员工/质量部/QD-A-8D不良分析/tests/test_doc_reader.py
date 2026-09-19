"""doc_reader 单测 — 段落解析与文本提取。"""
from qda_prefill.doc_reader import _parse_sections, DocumentSections, extract_scene_checkbox


def test_parse_d_sections_from_text():
    text = (
        "8D报告 - 案例编号: 8D-2026-001\n"
        "D2: 某产品在高温下异常重启，客户反馈。\n"
        "具体现象：偶发，频率约3次/百小时。\n"
        "D3: 紧急隔离现场库存。\n"
        "D4: 根本原因：电容降额不足。\n"
    )
    doc = _parse_sections(text, "test.docx")
    assert "D2" in doc.sections
    assert "电容降额不足" in doc.sections.get("D4", "")
    assert doc.full_text == text


def test_parse_returns_header_text():
    text = "X" * 600
    doc = _parse_sections(text, "test.docx")
    assert len(doc.header_text) == 500


def test_parse_no_sections_returns_empty_sections():
    text = "没有任何段落标头的普通文本内容。"
    doc = _parse_sections(text, "test.docx")
    assert doc.sections == {}
    assert doc.full_text == text


def test_read_docx_extracts_sections(synthetic_docx):
    from qda_prefill.doc_reader import read
    doc = read(synthetic_docx)
    assert isinstance(doc, DocumentSections)
    assert "D2" in doc.sections
    assert "D4" in doc.sections
    assert "D7" in doc.sections
    assert "D8" in doc.sections


# ── 场景勾选行（J7：PPT 勾选态，不是 Word w14:checkbox）──────────────────────

def test_scene_checkbox_reads_rnd_checked():
    r = extract_scene_checkbox("场景（必选）：☐ 制造 ☑ 研发")
    assert r is not None and r.scene == "研发" and r.ambiguous is False


def test_scene_checkbox_reads_manufacturing_checked():
    r = extract_scene_checkbox("场景（必选）：☑ 制造 ☐ 研发")
    assert r is not None and r.scene == "制造" and r.ambiguous is False


def test_scene_checkbox_absent_row_returns_none():
    assert extract_scene_checkbox("D2: 问题描述\n正常段落文本，无场景勾选行。") is None


def test_scene_checkbox_none_checked_is_ambiguous():
    r = extract_scene_checkbox("场景（必选）：☐ 制造 ☐ 研发")
    assert r is not None and r.scene is None and r.ambiguous is True


def test_scene_checkbox_both_checked_is_ambiguous():
    r = extract_scene_checkbox("场景（必选）：☑ 制造 ☑ 研发")
    assert r is not None and r.scene is None and r.ambiguous is True


def test_scene_checkbox_preserves_trailing_annotation():
    """验收集样本 3「☑ 制造 ※」——※ 未解释，本方不猜；raw_line 原样保留供校准报告单列。"""
    r = extract_scene_checkbox("场景（必选）：☑ 制造 ☐ 研发 ※")
    assert r is not None and r.scene == "制造" and r.ambiguous is False
    assert "※" in r.raw_line


def test_scene_checkbox_wired_into_document_sections():
    text = "D2: 问题描述\n场景（必选）：☐ 制造 ☑ 研发\n产品在高温下异常。\nD3: 临时对策\n隔离库存。\n"
    doc = _parse_sections(text, "test.pptx")
    assert doc.scene_checkbox is not None and doc.scene_checkbox.scene == "研发"


def test_synthetic_pptx_end_to_end_scene_checkbox(tmp_path):
    """搭一份最小合成 pptx（非真实 8D 内容），验证 read() 端到端读出 PPT 勾选态。"""
    from pptx import Presentation  # type: ignore
    from qda_prefill.doc_reader import read

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 9144000, 1000000)
    tf = box.text_frame
    tf.text = "D2: 问题描述"
    p2 = tf.add_paragraph()
    p2.text = "场景（必选）：☐ 制造 ☑ 研发"
    p3 = tf.add_paragraph()
    p3.text = "合成样本正文，用于校验解析链路，不含真实 8D 内容。"
    out = tmp_path / "synthetic_8d.pptx"
    prs.save(str(out))

    doc = read(out)
    assert doc.scene_checkbox is not None
    assert doc.scene_checkbox.scene == "研发"
    assert doc.scene_checkbox.ambiguous is False
